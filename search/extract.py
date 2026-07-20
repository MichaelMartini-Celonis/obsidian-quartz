"""Per-filetype text and lightweight bibliographic extraction.

Returns an :class:`ExtractedDoc` with page-anchored text so chunks can carry a
page number. Bibliographic parsing is best-effort; low-confidence results are
flagged via ``needs_review`` for later reconciliation.
"""

from __future__ import annotations

import json
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from . import config

BOILERPLATE_TITLES = {
    "untitled",
    "document",
    "microsoft word - document",
    "presentation",
    "contents lists available at sciencedirect",
}


@dataclass
class ExtractedDoc:
    title: str | None = None
    authors: list[str] = field(default_factory=list)
    year: int | None = None
    venue: str | None = None
    abstract: str | None = None
    pages: list[tuple[int, str]] = field(default_factory=list)  # (page_no, text)
    n_pages: int = 0
    extraction_confidence: float = 0.5
    needs_review: bool = False


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "").strip())


def _is_boilerplate(text: str) -> bool:
    low = _clean(text).lower()
    if not low or low in BOILERPLATE_TITLES:
        return True
    if low.startswith("microsoft word"):
        return True
    if re.match(r"^\d+(st|nd|rd|th)\s", low):
        return True
    return False


def _guess_year(text: str) -> int | None:
    years = [int(y) for y in re.findall(r"\b(19[7-9]\d|20[0-4]\d)\b", text[:4000])]
    return max(years) if years else None


def _extract_pdf(path: Path) -> ExtractedDoc:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: list[tuple[int, str]] = []
    # Stop extracting once we have well past the per-doc cap: extraction (not just
    # embedding) dominates cost on multi-thousand-page books, so cut it early.
    cap = int(config.MAX_CHARS_PER_DOC * 1.2) if config.MAX_CHARS_PER_DOC else 0
    total = 0
    for i, page in enumerate(reader.pages):
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt.strip():
            pages.append((i + 1, txt))
            total += len(txt)
        if cap and total >= cap:
            break

    title = None
    authors: list[str] = []
    meta = reader.metadata
    if meta and meta.title and not _is_boilerplate(str(meta.title)):
        title = _clean(str(meta.title))
    if meta and meta.author:
        authors = [_clean(a) for a in str(meta.author).split(",") if _clean(a)]

    head = "\n".join(t for _, t in pages[:2])
    lines = [_clean(l) for l in head.splitlines() if _clean(l)]
    if not title:
        for line in lines[:15]:
            if _is_boilerplate(line):
                continue
            if len(line.split()) >= 3 and "@" not in line:
                title = line
                break

    doc = ExtractedDoc(
        title=title or path.stem,
        authors=authors,
        year=_guess_year(head),
        pages=pages,
        n_pages=len(reader.pages),
    )
    doc.needs_review = title is None or not authors or _is_boilerplate(title or "")
    doc.extraction_confidence = 0.8 if (title and authors) else 0.4
    return doc


def _extract_docx(path: Path) -> ExtractedDoc:
    from docx import Document

    d = Document(str(path))
    paras = [p.text for p in d.paragraphs if p.text and p.text.strip()]
    text = "\n\n".join(paras)
    title = None
    if d.core_properties.title and not _is_boilerplate(d.core_properties.title):
        title = _clean(d.core_properties.title)
    if not title and paras:
        title = _clean(paras[0])
    authors = [_clean(d.core_properties.author)] if d.core_properties.author else []
    return ExtractedDoc(
        title=title or path.stem,
        authors=authors,
        year=_guess_year(text),
        pages=[(1, text)] if text else [],
        n_pages=1,
        extraction_confidence=0.6,
        needs_review=not title,
    )


def _extract_pptx(path: Path) -> ExtractedDoc:
    try:
        from pptx import Presentation
    except Exception:
        return ExtractedDoc(title=path.stem, pages=[], n_pages=0, needs_review=True,
                            extraction_confidence=0.1)
    prs = Presentation(str(path))
    pages = []
    title = None
    for i, slide in enumerate(prs.slides):
        texts = [sh.text for sh in slide.shapes if hasattr(sh, "text") and sh.text.strip()]
        if texts:
            if title is None:
                title = _clean(texts[0].splitlines()[0])
            pages.append((i + 1, "\n".join(texts)))
    return ExtractedDoc(title=title or path.stem, pages=pages, n_pages=len(pages),
                        extraction_confidence=0.5)


def _extract_ipynb(path: Path) -> ExtractedDoc:
    nb = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    parts = []
    title = None
    for cell in nb.get("cells", []):
        src = "".join(cell.get("source", []))
        if not src.strip():
            continue
        if cell.get("cell_type") == "markdown":
            if title is None:
                for line in src.splitlines():
                    if line.strip().startswith("#"):
                        title = _clean(line.lstrip("#"))
                        break
            parts.append(src)
        else:
            parts.append(src)
    text = "\n\n".join(parts)
    return ExtractedDoc(title=title or path.stem.replace("_", " "), pages=[(1, text)] if text else [],
                        n_pages=1, extraction_confidence=0.5)


def _extract_text(path: Path) -> ExtractedDoc:
    text = path.read_text(encoding="utf-8", errors="ignore")
    title = None
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("#"):
            title = _clean(s.lstrip("#"))
            break
        if s:
            title = _clean(s)
            break
    return ExtractedDoc(title=title or path.stem, pages=[(1, text)] if text.strip() else [],
                        n_pages=1, year=_guess_year(text), extraction_confidence=0.5)


def _extract_bpmn(path: Path) -> ExtractedDoc:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    names = re.findall(r'name="([^"]+)"', raw)
    proc = re.search(r'<(?:bpmn:)?process[^>]*name="([^"]+)"', raw)
    text = "\n".join(dict.fromkeys(names))  # unique, order-preserving
    return ExtractedDoc(title=_clean(proc.group(1)) if proc else path.stem,
                        pages=[(1, text)] if text else [], n_pages=1,
                        extraction_confidence=0.4)


def _extract_epub(path: Path) -> ExtractedDoc:
    title = path.stem.replace("_", " ")
    authors: list[str] = []
    text_parts: list[str] = []
    try:
        with zipfile.ZipFile(path) as z:
            opf = next((n for n in z.namelist() if n.endswith(".opf")), None)
            if opf:
                meta = z.read(opf).decode("utf-8", errors="ignore")
                m = re.search(r"<dc:title[^>]*>([^<]+)", meta)
                if m:
                    title = _clean(m.group(1))
                c = re.search(r"<dc:creator[^>]*>([^<]+)", meta)
                if c:
                    authors = [_clean(c.group(1))]
            for name in z.namelist():
                if name.endswith((".xhtml", ".html", ".htm")):
                    html = z.read(name).decode("utf-8", errors="ignore")
                    text_parts.append(re.sub(r"<[^>]+>", " ", html))
    except Exception:
        pass
    text = _clean("\n".join(text_parts))
    return ExtractedDoc(title=title, authors=authors, pages=[(1, text)] if text else [],
                        n_pages=1, extraction_confidence=0.4)


_DISPATCH = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".ipynb": _extract_ipynb,
    ".md": _extract_text,
    ".txt": _extract_text,
    ".bpmn": _extract_bpmn,
    ".epub": _extract_epub,
}


def extract_document(path: Path) -> ExtractedDoc:
    fn = _DISPATCH.get(path.suffix.lower())
    if fn is None:
        return ExtractedDoc(title=path.stem, pages=[], n_pages=0, needs_review=True,
                            extraction_confidence=0.0)
    try:
        return fn(path)
    except Exception as exc:  # extraction should never crash the pipeline
        return ExtractedDoc(title=path.stem, pages=[], n_pages=0, needs_review=True,
                            extraction_confidence=0.0, abstract=f"[extract error: {exc}]")
